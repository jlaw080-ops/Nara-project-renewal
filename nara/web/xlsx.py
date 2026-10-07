"""사업 조회를 출력 양식(15칸) 그대로 엑셀로 만든다."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from nara.web.data import PrintRow

HEADERS = (
    "수요기관", "공고명", "주소", "공고일(예정일)", "낙찰일(개찰일)", "착공일", "준공(예정)일",
    "담당부서", "부서장", "직위", "전화번호", "낙찰업체 설계사무소", "설치계획내용",
    "신재생에너지원별 예상가", "진행현황",
)  # fmt: skip
# print.html colgroup 비율(%)을 엑셀 열 너비로
WIDTHS = (6, 10, 9, 5, 5, 5, 5, 6, 4, 4, 6, 8, 9, 8, 10)
WIDTH_SCALE = 2.2
LEFT_COLUMNS = (12, 13)  # 설치계획내용·예상가는 여러 줄이라 왼쪽 정렬
HEAD_BG = "1F3864"
LINE = "9A9A9A"
FONT = "맑은 고딕"


def _costs(r: PrintRow) -> str:
    lines = list(r.costs)
    if r.costs:
        lines.append(f"합계 {r.total:,}원")
    if r.unpriced:
        lines.append(f"단가 없음: {', '.join(r.unpriced)}")
    return "\n".join(lines)


def _values(r: PrintRow) -> list[str]:
    return [
        r.org_name, r.name, r.address or "", r.notice_date or "", r.open_date or "",
        r.start_date or "", r.end_date or "", r.dept or "", r.head_name or "",
        r.head_position or "", r.dept_tel or "", r.winner or "", "\n".join(r.plan),
        _costs(r), r.status,
    ]  # fmt: skip


def workbook(rows: list[PrintRow], note: str = "") -> bytes:
    """머리글 고정·자동 필터·A4 가로 인쇄(머리글 반복). note가 있으면 표 아래에 적는다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "사업 조회"
    side = Side(style="thin", color=LINE)
    box = Border(side, side, side, side)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEAD_BG)
        cell.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = box
    for r in rows:
        ws.append(_values(r))
    for row in ws.iter_rows(min_row=2):
        for i, cell in enumerate(row):
            cell.font = Font(name=FONT, size=9)
            cell.border = box
            align = "left" if i in LEFT_COLUMNS else "center"
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width * WIDTH_SCALE
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if note:
        ws.append([])
        ws.append([note])
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
